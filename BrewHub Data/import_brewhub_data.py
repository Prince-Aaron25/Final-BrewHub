"""One-time Excel importer for BrewHub.

Expected workbook sheets/headers are the ones in BrewHub_Products_Recipes_Inventory.xlsx.
Run from your BrewHub project environment after installing openpyxl and mysql-connector-python.

IMPORTANT: Back up brewhub_db first. The script does not delete existing records or import orders/payments.
"""
from pathlib import Path
import re
import shutil
import sys
from decimal import Decimal, InvalidOperation

import mysql.connector
from openpyxl import load_workbook

# ---- EDIT THESE SETTINGS ----
DB_CONFIG = {
    "host": "localhost",
    "port": 3307,
    "user": "root",                 # Use your actual MySQL user
    "password": "Aaronsacang0825", # Replace locally; never share this file with password filled in
    "database": "brewhub_db",
}
WORKBOOK = Path(__file__).with_name("BrewHub_Products_Recipes_Inventory.xlsx")
IMAGE_FOLDER = Path(__file__).with_name("product_images")  # Put product image files here
# The saved path is relative to the BrewHub project folder. Change if your app expects another format.
IMAGE_DB_PREFIX = "product_images"
# -----------------------------


def clean(value):
    return " ".join(str(value).strip().split()) if value is not None else ""


def norm(value):
    return re.sub(r"[^a-z0-9]+", "", clean(value).lower())


def decimal_value(value, label):
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError(f"Invalid numeric value for {label}: {value!r}")


def read_sheet(ws, expected_headers):
    """Find the header row instead of assuming the sheet starts at row 1."""
    header_row = None
    headers = None
    expected = {norm(x) for x in expected_headers}
    for row_num, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = [clean(v) for v in row]
        found = {norm(v) for v in values if v}
        if expected.issubset(found):
            header_row, headers = row_num, values
            break
    if header_row is None:
        raise ValueError(f"Could not find expected headers in sheet {ws.title!r}: {expected_headers}")
    index = {norm(name): i for i, name in enumerate(headers) if name}
    records = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        record = {key: (row[index[norm(header)]] if index[norm(header)] < len(row) else None)
                  for key, header in ((norm(h), h) for h in expected_headers)}
        if any(v is not None and clean(v) for v in record.values()):
            records.append(record)
    return records


def get(rec, header):
    return rec.get(norm(header))


def find_image(product_name):
    if not IMAGE_FOLDER.exists():
        return None
    supported = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
    target = norm(product_name)
    # Match image filename to product name, ignoring spaces/punctuation/case.
    for f in IMAGE_FOLDER.iterdir():
        if f.is_file() and f.suffix.lower() in supported and norm(f.stem) == target:
            return f
    return None


def main():
    if not WORKBOOK.exists():
        sys.exit(f"Workbook not found: {WORKBOOK}\nPlace the Excel file beside this script or edit WORKBOOK.")
    if DB_CONFIG["password"] == "YOUR_MYSQL_PASSWORD":
        sys.exit("Edit DB_CONFIG['password'] in this script before running it.")

    wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
    products = read_sheet(wb["Products"], ["Product name", "Category", "Product group", "Selling unit", "Price (PHP)"])
    recipes = read_sheet(wb["Recipes"], ["Product name", "Ingredient name", "Quantity required", "Unit"])
    inventory = read_sheet(wb["Inventory"], ["Ingredient name", "Unit", "Actual stock quantity", "Reorder level"])

    print(f"Workbook: {WORKBOOK.name}")
    print(f"Products: {len(products)} | Recipe rows: {len(recipes)} | Inventory rows: {len(inventory)}")
    print("Image folder:", IMAGE_FOLDER)
    print("\nThis importer will add/update catalog data only; it will not create orders or payments.")
    confirm = input("Type IMPORT to proceed (anything else cancels): ").strip()
    if confirm != "IMPORT":
        print("Cancelled. No database changes made.")
        return

    conn = mysql.connector.connect(**DB_CONFIG)
    cursor = conn.cursor(dictionary=True)
    counts = {"categories_added": 0, "products_added": 0, "products_updated": 0,
              "ingredients_added": 0, "ingredients_updated": 0, "recipes_added": 0,
              "recipes_updated": 0, "images_matched": 0, "images_missing": 0,
              "recipe_skipped": 0, "warnings": []}
    try:
        conn.start_transaction()

        # Create categories first.
        category_names = sorted({clean(get(r, "Category")) for r in products if clean(get(r, "Category"))})
        for category_name in category_names:
            cursor.execute("SELECT category_id FROM categories WHERE LOWER(category_name)=LOWER(%s) LIMIT 1", (category_name,))
            if not cursor.fetchone():
                cursor.execute("INSERT INTO categories (category_name) VALUES (%s)", (category_name,))
                counts["categories_added"] += 1
        cursor.execute("SELECT category_id, category_name FROM categories")
        category_map = {norm(r["category_name"]): r["category_id"] for r in cursor.fetchall()}

        product_map = {}

        for r in products:
            name = clean(get(r, "Product name"))
            price_raw = get(r, "Price (PHP)")

            # Skip blank rows and footer/notes rows (no price = not a product)
            if not name or price_raw is None or not clean(price_raw):
                continue

            price = decimal_value(price_raw, f"price for {name}")
            category_id = category_map.get(norm(get(r, "Category")))
            if category_id is None:
                counts["warnings"].append(f"Product skipped: {name} (category {get(r, 'Category')!r} not found).")
                continue

            selling_unit = clean(get(r, "Selling unit"))
            group = clean(get(r, "Product group"))
            notes = clean(get(r, "Recipe notes"))
            description_parts = [x for x in (group, selling_unit, notes) if x]
            description = " | ".join(description_parts)[:255] or None
            image = find_image(name)
            image_path = f"{IMAGE_DB_PREFIX}/{image.name}" if image else None
            if image:
                counts["images_matched"] += 1
            else:
                counts["images_missing"] += 1

            cursor.execute("SELECT product_id, image_path FROM products WHERE LOWER(product_name)=LOWER(%s) ORDER BY product_id LIMIT 1", (name,))
            existing = cursor.fetchone()
            if existing:
                # Preserve an existing image path when no matching image file is supplied.
                final_image = image_path if image_path else existing["image_path"]
                cursor.execute("UPDATE products SET category_id=%s, description=%s, price=%s, image_path=%s, availability=COALESCE(availability,1) WHERE product_id=%s",
                               (category_id, description, price, final_image, existing["product_id"]))
                product_id = existing["product_id"]
                counts["products_updated"] += 1
            else:
                cursor.execute("INSERT INTO products (category_id, product_name, description, price, image_path, availability) VALUES (%s,%s,%s,%s,%s,1)",
                               (category_id, name, description, price, image_path))
                product_id = cursor.lastrowid
                counts["products_added"] += 1
            product_map[norm(name)] = product_id

        # Ingredients: match by normalized name; never invent stock levels.
        inventory_by_name = {norm(clean(get(r, "Ingredient name"))): r for r in inventory if clean(get(r, "Ingredient name"))}
        cursor.execute("SELECT ingredient_id, ingredient_name, quantity, unit, reorder_level FROM ingredients")
        existing_ingredients = {norm(r["ingredient_name"]): r for r in cursor.fetchall()}
        ingredient_map = {}
        all_recipe_ingredients = {norm(clean(get(r, "Ingredient name"))): clean(get(r, "Ingredient name"))
                                  for r in recipes if clean(get(r, "Ingredient name"))}
        all_names = dict(all_recipe_ingredients)
        for key, inv in inventory_by_name.items():
            all_names.setdefault(key, clean(get(inv, "Ingredient name")))

        for key, name in all_names.items():
            inv = inventory_by_name.get(key)
            unit = clean(get(inv, "Unit")) if inv else ""
            if not unit:
                # Fall back to the unit stated on a recipe row for this ingredient.
                unit = next((clean(get(rr, "Unit")) for rr in recipes
                             if norm(clean(get(rr, "Ingredient name"))) == key and clean(get(rr, "Unit"))), "")
            if not unit:
                counts["warnings"].append(f"No unit found for ingredient {name}; skipped.")
                continue
            stock = get(inv, "Actual stock quantity") if inv else None
            reorder = get(inv, "Reorder level") if inv else None
            stock_val = decimal_value(stock, f"stock for {name}") if stock is not None and clean(stock) else None
            reorder_val = decimal_value(reorder, f"reorder level for {name}") if reorder is not None and clean(reorder) else None
            old = existing_ingredients.get(key)
            if old:
                # Keep current stock/reorder values when Excel cells are blank.
                new_stock = stock_val if stock_val is not None else old["quantity"]
                new_reorder = reorder_val if reorder_val is not None else old["reorder_level"]
                cursor.execute("UPDATE ingredients SET unit=%s, quantity=%s, reorder_level=%s WHERE ingredient_id=%s",
                               (unit, new_stock, new_reorder, old["ingredient_id"]))
                ingredient_id = old["ingredient_id"]
                counts["ingredients_updated"] += 1
            else:
                cursor.execute("INSERT INTO ingredients (ingredient_name, quantity, unit, reorder_level) VALUES (%s,%s,%s,%s)",
                               (name, stock_val, unit, reorder_val))
                ingredient_id = cursor.lastrowid
                counts["ingredients_added"] += 1
            ingredient_map[key] = ingredient_id

        # Recipe relationship rows: one row per product/ingredient. Check units to prevent silent mismatches.
        for r in recipes:
            product_name = clean(get(r, "Product name"))
            ingredient_name = clean(get(r, "Ingredient name"))
            if not product_name or not ingredient_name:
                continue
            pid = product_map.get(norm(product_name))
            iid = ingredient_map.get(norm(ingredient_name))
            if not pid or not iid:
                counts["recipe_skipped"] += 1
                counts["warnings"].append(f"Recipe skipped: {product_name} / {ingredient_name} (product or ingredient not mapped).")
                continue
            
            if get(r, "Quantity required") is None or not clean(get(r, "Quantity required")):
                continue
            
            qty = decimal_value(get(r, "Quantity required"), f"recipe quantity {product_name}/{ingredient_name}")
            recipe_unit = clean(get(r, "Unit"))
            cursor.execute("SELECT unit FROM ingredients WHERE ingredient_id=%s", (iid,))
            ingredient_unit = cursor.fetchone()["unit"]
            if recipe_unit and norm(recipe_unit) != norm(ingredient_unit):
                counts["recipe_skipped"] += 1
                counts["warnings"].append(f"Unit mismatch: {product_name} needs {ingredient_name} in {recipe_unit}, inventory unit is {ingredient_unit}; recipe row skipped.")
                continue
            cursor.execute("SELECT quantity_required FROM product_ingredients WHERE product_id=%s AND ingredient_id=%s", (pid, iid))
            existing_recipe = cursor.fetchone()
            if existing_recipe:
                cursor.execute("UPDATE product_ingredients SET quantity_required=%s WHERE product_id=%s AND ingredient_id=%s", (qty, pid, iid))
                counts["recipes_updated"] += 1
            else:
                cursor.execute("INSERT INTO product_ingredients (product_id, ingredient_id, quantity_required) VALUES (%s,%s,%s)", (pid, iid, qty))
                counts["recipes_added"] += 1

        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cursor.close()
        conn.close()
        wb.close()

    print("\nIMPORT FINISHED")
    for key, value in counts.items():
        if key != "warnings":
            print(f"{key}: {value}")
    if counts["warnings"]:
        print("\nWARNINGS:")
        for warning in counts["warnings"]:
            print("-", warning)
    print("\nCheck products.image_path against the path format used by your GUI.")
    print("No ingredient stock was invented. Blank Excel stock cells remain NULL for new ingredients and preserve existing stock for existing ingredients.")


if __name__ == "__main__":
    main()
